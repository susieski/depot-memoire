# -*- coding: utf-8 -*-
"""Mesure du filtrage lexical sur le corpus événementiel."""
import csv, os, re, sys, unicodedata, collections
from pathlib import Path
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lexique119 import LEXIQUE

DEPOT = Path(__file__).resolve().parent.parent
SOURCES = Path(os.environ.get("CORPUS_SOURCES", DEPOT.parent / "donnees_sources"))


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "").lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


MOTIFS = []
for typ, termes in LEXIQUE.items():
    for t in termes:
        noyau = re.escape(norm(t)).replace(r"\+", ".{0,3}").replace(r"\-", "[- ]")
        MOTIFS.append((typ, t, re.compile(r"(?<![a-z0-9])" + noyau + r"[a-z]{0,3}(?![a-z])")))
assert len(MOTIFS) == 119, f"{len(MOTIFS)} motifs au lieu de 119"


def trouver(texte):
    t = norm(texte)
    return [(typ, terme) for typ, terme, rx in MOTIFS if rx.search(t)]


def lire_csv(nom):
    with open(DEPOT / "data" / nom, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


ws = openpyxl.load_workbook(
    SOURCES / "FUSION annotations + accord.xlsx", read_only=True)["Fusion"]
textes_lex = {str(l[0]).strip(): (l[4] or "") for l in list(ws.iter_rows(values_only=True))[1:] if l[0]}
lex = lire_csv("corpus_lexical_etiquettes.csv")

vus, exacts, n = 0, 0, 0
for r in lex:
    txt = textes_lex.get(r["item"].strip(), "")
    if not txt:
        continue
    n += 1
    hits = trouver(txt)
    vus += bool(hits)
    exacts += any(norm(te) == norm(r["terme_declencheur"].replace("_", " ")) for _, te in hits)

print("contrôle — corpus lexical")
print(f"  un terme du lexique retrouvé : {vus}/{n} = {vus/n:.1%}")
print(f"  le déclencheur lui-même      : {exacts}/{n} = {exacts/n:.1%}")

textes = {}
for f in ("echantillon_final_100_avant.xlsx", "echantillon_final_100_apres.xlsx"):
    w = openpyxl.load_workbook(SOURCES / f, read_only=True).worksheets[0]
    for r in list(w.iter_rows(values_only=True))[1:]:
        if r[1]:
            textes[str(r[1]).strip()] = r[3] or ""

evt = lire_csv("corpus_evenementiel_etiquettes.csv")
manquants = [r["item"] for r in evt if r["id_tweet"].strip() not in textes]
if manquants:
    sys.exit(f"texte introuvable : items {manquants}")

tab = collections.Counter()
for r in evt:
    hits = trouver(textes[r["id_tweet"].strip()])
    r["terme_lexique"] = "1" if hits else "0"
    r["termes_trouves"] = ";".join(sorted({t for _, t in hits}))
    tab[(r["reference"], bool(hits))] += 1

chemin = DEPOT / "data" / "corpus_evenementiel_etiquettes.csv"
with open(chemin, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(evt[0].keys()))
    w.writeheader()
    w.writerows(evt)

print(f"\nécrit : {chemin.name}")
print("\nlexique appliqué au corpus événementiel")
for c, nom in (("0", "non sexiste"), ("1", "hostile"), ("2", "implicite")):
    tot = tab[(c, True)] + tab[(c, False)]
    print(f"  classe {c} ({nom:11s}) n = {tot:3d} : capté {tab[(c,True)]:3d}   manqué {tab[(c,False)]:3d}")
capte = tab[("1", True)] + tab[("2", True)]
sexistes = capte + tab[("1", False)] + tab[("2", False)]
print(f"\n  rappel classes 1 et 2 : {capte}/{sexistes} = {capte/sexistes:.1%}")
print(f"  rappel classe 2       : {tab[('2',True)]}/{tab[('2',True)]+tab[('2',False)]}")
print("\n  tweets captés :")
for r in evt:
    if r["terme_lexique"] == "1":
        print(f"    item {r['item']:>3}  classe {r['reference']}  ({r['moment']})  terme : {r['termes_trouves']}")
