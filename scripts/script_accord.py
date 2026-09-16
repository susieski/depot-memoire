# -*- coding: utf-8 -*-
"""Calcul des accords — corpus lexical."""
import csv
import itertools
import os
from collections import Counter
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
DEPOT = BASE.parent
SOURCES = Path(os.environ.get("CORPUS_SOURCES", DEPOT.parent / "donnees_sources"))
SORTIE = SOURCES / "FUSION annotations + accord.xlsx"
SORTIE_METRIQUES = DEPOT / "resultats" / "metriques_corpus_lexical.csv"
DECISIONS_FINALES = DEPOT / "data" / "decisions_finales.csv"

FICHIERS = {
    "H1":     SOURCES / "Annotation corpus lexical - Annotatrice 1 h 1.xlsx",
    "H2":     SOURCES / "Annotation corpus lexical - Annotatrice 2 h.xlsx",
    "H3":     SOURCES / "Annotation corpus lexical - Annotateur 3-h.xlsx",
    "Grok":   SOURCES / "annotations_grok_lexical_passe1.xlsx",
    "Gemini": SOURCES / "annotations_gemini_lexical_passe1.xlsx",
}
FICHIER_TEXTES = SOURCES / "TEXTES échantillon 200 (ne pas utiliser les colonnes annotation).xlsx"
CLE = SOURCES / "CLE échantillon - ne pas diffuser.xlsx"
HUMAINS = ["H1", "H2", "H3"]

if not SOURCES.exists():
    raise SystemExit(f"Classeurs sources introuvables : {SOURCES}\nIndiquez leur emplacement dans CORPUS_SOURCES.")


def lire_annotation(f):
    xl = pd.ExcelFile(f)
    feuille = "Annotation" if "Annotation" in xl.sheet_names else xl.sheet_names[0]
    df = xl.parse(feuille)
    df.columns = [str(c).strip() for c in df.columns]
    ren = {}
    for c in df.columns:
        cl = c.lower()
        if cl.startswith("n°") or cl == "n": ren[c] = "n"
        elif cl.startswith("classe"): ren[c] = "classe"
        elif cl.startswith("incertain"): ren[c] = "incertain"
        elif cl.startswith("code"): ren[c] = "codes"
        elif cl.startswith("tweet_parent"): ren[c] = "parent"
        elif cl == "tweet": ren[c] = "tweet"
        elif cl == "id_tweet": ren[c] = "id_tweet"
    df = df.rename(columns=ren)
    df = df[[c for c in ("n","id_tweet","tweet","parent","classe","incertain","codes") if c in df.columns]]
    df = df.dropna(subset=["n"]).copy()
    df["n"] = df["n"].astype(int)
    df["classe"] = pd.to_numeric(df["classe"], errors="coerce")
    df.loc[~df["classe"].isin([0,1,2]), "classe"] = np.nan
    return df.set_index("n")


def kappa(y1, y2, labels=(0,1,2)):
    po = (y1 == y2).mean()
    pe = sum((y1 == c).mean() * (y2 == c).mean() for c in labels)
    return (po - pe) / (1 - pe)

def kappa_ic(y1, y2, B=2000, seed=20261206):
    rng = np.random.default_rng(seed)
    n = len(y1)
    ks = [kappa(y1[i], y2[i]) for i in (rng.integers(0, n, n) for _ in range(B))]
    return float(np.percentile(ks, 2.5)), float(np.percentile(ks, 97.5))

def alpha_krippendorff(colonnes):
    units = [r.dropna().astype(int).tolist() for _, r in colonnes.iterrows()]
    units = [u for u in units if len(u) >= 2]
    Do_num = Do_den = 0
    for u in units:
        m = len(u)
        Do_num += sum(a != b for a, b in itertools.permutations(u, 2)) / (m - 1)
        Do_den += m
    Do = Do_num / Do_den
    allv = [v for u in units for v in u]
    n = len(allv)
    c = Counter(allv)
    De = sum(c[a] * c[b] for a in c for b in c if a != b) / (n * (n - 1))
    return 1 - Do / De

def prf(y_vrai, y_pred, c):
    vp = ((y_pred == c) & (y_vrai == c)).sum()
    fp = ((y_pred == c) & (y_vrai != c)).sum()
    fn = ((y_pred != c) & (y_vrai == c)).sum()
    P = vp / (vp + fp) if vp + fp else 0.0
    R = vp / (vp + fn) if vp + fn else 0.0
    F = 2 * P * R / (P + R) if P + R else 0.0
    return P, R, F


annos = {nom: lire_annotation(f) for nom, f in FICHIERS.items()}
cle = pd.read_excel(CLE)
cle.columns = [str(c).strip() for c in cle.columns]
cle = cle.rename(columns={"n° (ordre fichier)": "n"}).set_index("n")

for nom, df in annos.items():
    if "id_tweet" in df.columns:
        commun = df.index.intersection(cle.index)
        ok = (df.loc[commun, "id_tweet"].astype(str) == cle.loc[commun, "id_tweet"].astype(str)).all()
        assert ok, f"ids non alignés : {nom}"

ref = lire_annotation(FICHIER_TEXTES)
fusion = pd.DataFrame(index=cle.index)
fusion["id_tweet"] = cle["id_tweet"].astype(str)
fusion["type"] = cle["strate"] if "strate" in cle.columns else cle["type"]
fusion["terme déclencheur"] = cle["terme déclencheur"]
fusion["tweet"] = ref["tweet"]
fusion["tweet_parent"] = ref["parent"]
for nom in ("H1", "H2", "H3", "Grok", "Gemini"):
    fusion[nom] = annos[nom]["classe"]
    fusion[f"incertain_{nom}"] = annos[nom]["incertain"].astype(str).str.strip().str.lower().eq("oui")

def config_ligne(r):
    votes = [v for v in (r["H1"], r["H2"], r["H3"]) if pd.notna(v)]
    if len(votes) < 2:
        return pd.Series([np.nan, "insuffisant", True])
    cnt = Counter(int(v) for v in votes)
    (c1, n1), *reste = cnt.most_common()
    if len(cnt) == 1:
        return pd.Series([c1, "unanimité" if len(votes) == 3 else "accord (2 voix)", False])
    if n1 > reste[0][1]:
        return pd.Series([c1, "majorité 2-1", False])
    return pd.Series([np.nan, "éclaté / égalité", True])

fusion[["majorité_humaine", "configuration", "à_adjudiquer"]] = fusion.apply(config_ligne, axis=1)


A = fusion[["H1", "H2", "H3", "Grok", "Gemini"]]
paires = []
for a, b in itertools.combinations(A.columns, 2):
    sub = A[[a, b]].dropna()
    y1, y2 = sub[a].astype(int).values, sub[b].astype(int).values
    lo, hi = kappa_ic(y1, y2)
    paires.append((f"{a} × {b}", len(sub), round((y1 == y2).mean(), 3),
                   round(kappa(y1, y2), 3), round(lo, 3), round(hi, 3)))

alphas = [("3 humains", round(alpha_krippendorff(A[HUMAINS]), 3)),
          ("3 humains + Grok", round(alpha_krippendorff(A[HUMAINS + ["Grok"]]), 3)),
          ("3 humains + Gemini", round(alpha_krippendorff(A[HUMAINS + ["Gemini"]]), 3)),
          ("les 5", round(alpha_krippendorff(A), 3))]

maj = fusion["majorité_humaine"]

def eval_llm(llm, index, reference=None):
    ref = maj if reference is None else reference
    sub = pd.concat([ref.rename("reference"), A[llm]], axis=1).loc[index].dropna()
    yv, yp = sub["reference"].astype(int).values, sub[llm].astype(int).values
    ligne = {"llm": llm, "n": len(sub), "accord": round((yv == yp).mean(), 3),
             "kappa": round(kappa(yv, yp), 3)}
    for c in (0, 1, 2):
        P, R, F = prf(yv, yp, c)
        ligne[f"P{c}"], ligne[f"R{c}"], ligne[f"F{c}"] = round(P, 2), round(R, 2), round(F, 2)
    return ligne

llm_eval = [eval_llm(llm, A.index) for llm in ("Grok", "Gemini")]
constant = fusion.index[maj.notna() & A["Grok"].notna() & A["Gemini"].notna()]
llm_eval_constant = [eval_llm(llm, constant) for llm in ("Grok", "Gemini")]

decisions_finales = pd.read_csv(DECISIONS_FINALES).set_index("item")
manquants = sorted(set(fusion.index[maj.isna()]) - set(decisions_finales.index))
if manquants:
    raise SystemExit(f"Items sans majorité ni décision finale : {manquants}")
reference_finale = maj.copy()
for item, ligne_decision in decisions_finales.iterrows():
    reference_finale.loc[item] = ligne_decision["decision"]
fusion["référence_finale"] = reference_finale
llm_eval_reference = [eval_llm(llm, A.index, reference_finale) for llm in ("Grok", "Gemini")]


ARIAL = Font(name="Arial", size=10)
GRAS = Font(name="Arial", size=10, bold=True)
TITRE = Font(name="Arial", size=12, bold=True)
GRIS = PatternFill("solid", fgColor="DDDDDD")
JAUNE = PatternFill("solid", fgColor="FFFF00")

def style_ligne(ws, ligne, gras=False, fond=None):
    for cell in ws[ligne]:
        cell.font = GRAS if gras else ARIAL
        if fond: cell.fill = fond

wb = Workbook()

ws = wb.active
ws.title = "Fusion"
cols = ["n°", "id_tweet", "type", "terme déclencheur", "tweet", "tweet_parent",
        "H1", "H2", "H3", "Grok", "Gemini",
        "incertain_H1", "incertain_H2", "incertain_H3", "incertain_Grok", "incertain_Gemini",
        "majorité_humaine", "configuration", "à_adjudiquer"]
ws.append(cols)
style_ligne(ws, 1, gras=True, fond=GRIS)
for n, r in fusion.iterrows():
    ws.append([n, r["id_tweet"], r["type"], r["terme déclencheur"], r["tweet"], r["tweet_parent"]]
              + [int(r[c]) if pd.notna(r[c]) else None for c in ("H1","H2","H3","Grok","Gemini")]
              + [("oui" if r[f"incertain_{c}"] else "") for c in ("H1","H2","H3","Grok","Gemini")]
              + [int(r["majorité_humaine"]) if pd.notna(r["majorité_humaine"]) else None,
                 r["configuration"], "OUI" if r["à_adjudiquer"] else ""])
    style_ligne(ws, ws.max_row)
ws.freeze_panes = "B2"
for i, w in enumerate([5,20,22,16,60,45,5,5,5,6,7,6,6,6,6,6,9,14,12], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws = wb.create_sheet("Métriques")
def ecrire(valeurs, gras=False, fond=None):
    ws.append(valeurs); style_ligne(ws, ws.max_row, gras=gras, fond=fond)

ecrire([f"Métriques d'accord — corpus lexical (généré le {date.today().isoformat()} par script_accord.py)"], gras=True)
ws["A1"].font = TITRE
gem_manq = int(A["Gemini"].isna().sum())
if gem_manq:
    ecrire([f"Annotation Gemini incomplète : {200-gem_manq}/200."], gras=True, fond=JAUNE)
ecrire([])
ecrire(["Distributions (n items par classe)"], gras=True, fond=GRIS)
ecrire(["annotateur", "classe 0", "classe 1", "classe 2", "manquants", "incertains"], gras=True)
for c in ("H1","H2","H3","Grok","Gemini"):
    v = A[c].value_counts()
    ecrire([c, int(v.get(0,0)), int(v.get(1,0)), int(v.get(2,0)), int(A[c].isna().sum()),
            int(fusion[f"incertain_{c}"].sum())])
ecrire([])
ecrire(["Accords par paire (kappa de Cohen, IC 95 % bootstrap, graine 20261206, B=2000)"], gras=True, fond=GRIS)
ecrire(["paire", "n items", "accord brut", "kappa", "IC bas", "IC haut"], gras=True)
for p in paires: ecrire(list(p))
ecrire([])
ecrire(["Alpha de Krippendorff (nominal)"], gras=True, fond=GRIS)
for nom, a in alphas: ecrire([nom, a])
ecrire([])
ecrire(["Configurations humaines"], gras=True, fond=GRIS)
for cfg, n in fusion["configuration"].value_counts().items(): ecrire([cfg, int(n)])
ecrire([])
ecrire(["LLM vs majorité humaine (items avec majorité et réponse LLM)"], gras=True, fond=GRIS)
ecrire(["LLM", "n", "accord", "kappa", "P0", "R0", "F1_0", "P1", "R1", "F1_1", "P2", "R2", "F1_2"], gras=True)
for l in llm_eval:
    ecrire([l["llm"], l["n"], l["accord"], l["kappa"],
            l["P0"], l["R0"], l["F0"], l["P1"], l["R1"], l["F1"], l["P2"], l["R2"], l["F2"]])
ecrire([])
ecrire([f"SOUS-ÉCHANTILLON CONSTANT (n={len(constant)} : majorité présente + réponse des 2 LLM)"], gras=True, fond=GRIS)
ecrire(["LLM", "n", "accord", "kappa", "P0", "R0", "F1_0", "P1", "R1", "F1_1", "P2", "R2", "F1_2"], gras=True)
for l in llm_eval_constant:
    ecrire([l["llm"], l["n"], l["accord"], l["kappa"],
            l["P0"], l["R0"], l["F0"], l["P1"], l["R1"], l["F1"], l["P2"], l["R2"], l["F2"]])
ecrire([])
ecrire([f"RÉFÉRENCE FINALE (n={int(reference_finale.notna().sum())} : majorité humaine complétée par les "
        f"{len(decisions_finales)} décisions finales de data/decisions_finales.csv)"], gras=True, fond=GRIS)
ecrire(["LLM", "n", "accord", "kappa", "P0", "R0", "F1_0", "P1", "R1", "F1_1", "P2", "R2", "F1_2"], gras=True)
for l in llm_eval_reference:
    ecrire([l["llm"], l["n"], l["accord"], l["kappa"],
            l["P0"], l["R0"], l["F0"], l["P1"], l["R1"], l["F1"], l["P2"], l["R2"], l["F2"]])
ecrire([])
vref = reference_finale.value_counts()
ecrire(["Distribution de la référence finale", "classe 0", "classe 1", "classe 2"], gras=True)
ecrire(["n items", int(vref.get(0, 0)), int(vref.get(1, 0)), int(vref.get(2, 0))])
for i, w in enumerate([30,10,11,9,8,8,8,8,8,8,8,8,8], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws = wb.create_sheet("À adjudiquer")
ecrire2 = lambda v, **k: (ws.append(v), style_ligne(ws, ws.max_row, **k))
ws.append(["Colonnes à remplir : « décision » = 0, 1 ou 2 ; « justification » = règle du guide appliquée."])
style_ligne(ws, 1, gras=True, fond=JAUNE)
ws.append([])
entetes = ["n°", "tweet", "tweet_parent", "H1", "H2", "H3", "configuration", "décision", "justification"]
ws.append(entetes)
style_ligne(ws, 3, gras=True, fond=GRIS)
adj = fusion[fusion["à_adjudiquer"]]
for n, r in adj.iterrows():
    ws.append([n, r["tweet"], r["tweet_parent"]]
              + [int(r[c]) if pd.notna(r[c]) else None for c in ("H1","H2","H3")]
              + [r["configuration"], None, None])
    lig = ws.max_row
    style_ligne(ws, lig)
    ws.cell(row=lig, column=8).fill = JAUNE
    ws.cell(row=lig, column=9).fill = JAUNE
ws.freeze_panes = "A4"
for i, w in enumerate([5,70,50,5,5,5,16,10,45], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for ligne in ws.iter_rows(min_row=4):
    for cell in ligne:
        cell.alignment = Alignment(wrap_text=(cell.column in (2,3,9)), vertical="top")

wb.save(SORTIE)

SORTIE_METRIQUES.parent.mkdir(parents=True, exist_ok=True)
with open(SORTIE_METRIQUES, "w", encoding="utf-8-sig", newline="") as f:
    plume = csv.writer(f)
    for ligne in wb["Métriques"].iter_rows(values_only=True):
        valeurs = ["" if v is None else v for v in ligne]
        while valeurs and valeurs[-1] == "":
            valeurs.pop()
        plume.writerow(valeurs)

print(f"{SORTIE.name} | {SORTIE_METRIQUES.name} | {len(adj)} à adjudiquer")
print(f"  alpha 3 humains {alphas[0][1]} | " +
      ", ".join(f"{p[0]}={p[3]}" for p in paires if p[0].count('H') == 2))
print(f"  constant n={len(constant)} : " +
      " | ".join(f"{l['llm']} kappa={l['kappa']} accord={l['accord']}" for l in llm_eval_constant))
print(f"  référence n={int(reference_finale.notna().sum())} : " +
      " | ".join(f"{l['llm']} kappa={l['kappa']} accord={l['accord']}" for l in llm_eval_reference))
