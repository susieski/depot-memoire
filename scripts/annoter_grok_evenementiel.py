# -*- coding: utf-8 -*-
"""Annotation du corpus événementiel par Grok (API xAI). La consigne, la lecture du
corpus et l'analyse des réponses viennent de annoter_gemini_evenementiel.py.
"""

import json, os, sys, time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openai import OpenAI, APIError, APIStatusError, APIConnectionError

from annoter_gemini_evenementiel import (PROMPT, FICHIER_SOURCE,
                                         charger_tweets, parser_reponse)

MODELE = "grok-4.5"
TEMPERATURE = 0.0

DEPOT = Path(__file__).resolve().parent.parent
SOURCES = Path(os.environ.get("CORPUS_SOURCES", DEPOT.parent / "donnees_sources"))

PASSE = 1
LIGNES = None
MAX_APPELS = 250
PAUSE_ENTRE_APPELS = 1.0

FICHIER_SORTIE = SOURCES / f"annotations_grok_evenementiel_passe{PASSE}.xlsx"
FICHIER_JOURNAL = SOURCES / f"journal_grok_evenementiel_passe{PASSE}.json"

PRIX_ENTREE, PRIX_SORTIE = 2.00, 6.00


def preparer_sortie():
    a_reecrire = {}
    if os.path.exists(FICHIER_SORTIE):
        wb = openpyxl.load_workbook(FICHIER_SORTIE)
        ws = wb.active
        faits = set()
        for idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if r[0] is None:
                continue
            n = int(r[0])
            if r[2] is not None:
                faits.add(n)
            else:
                a_reecrire[n] = idx
        print(f"reprise : {len(faits)} faites, {len(a_reecrire)} à refaire")
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Annotation LLM"
        ws.append(["n°", "id_tweet", "classe", "incertain", "codes", "justification",
                   "refus", "modele_version", "horodatage_utc", "tokens_entree", "tokens_sortie"])
        faits = set()
    return wb, ws, faits, a_reecrire

def ecrire_ligne(ws, valeurs, a_reecrire):
    n = valeurs[0]
    if n in a_reecrire:
        for j, v in enumerate(valeurs, start=1):
            ws.cell(row=a_reecrire[n], column=j, value=v)
    else:
        ws.append(valeurs)

def appeler(client, texte, parent):
    prompt = PROMPT.replace("{{TEXTE}}", texte).replace("{{TEXTE_PARENT}}", parent)
    for tentative in (1, 2, 3):
        try:
            return client.chat.completions.create(
                model=MODELE,
                messages=[{"role": "user", "content": prompt}],
                temperature=TEMPERATURE,
                response_format={"type": "json_object"})
        except (APIConnectionError, APIStatusError, APIError) as e:
            code = getattr(e, "status_code", None)
            if (code in (429, 500, 502, 503) or isinstance(e, APIConnectionError)) and tentative < 3:
                attente = 15 * tentative
                print(f"  {code or 'réseau'} — nouvelle tentative dans {attente}s")
                time.sleep(attente)
            else:
                raise


def main():
    cle = os.environ.get("XAI_API_KEY", "")
    if not cle:
        sys.exit("Définissez XAI_API_KEY.")
    client = OpenAI(api_key=cle, base_url="https://api.x.ai/v1")

    tweets = charger_tweets()
    if LIGNES is not None:
        tweets = [t for t in tweets if t["n"] in set(LIGNES)]
    wb, ws, faits, a_reecrire = preparer_sortie()
    a_faire = [t for t in tweets if t["n"] not in faits]
    print(f"{len(a_faire)} tweets à annoter — {MODELE}, T={TEMPERATURE}, passe {PASSE}")

    journal = {"corpus": "evenementiel #MissFrance2026", "modele_demande": MODELE,
               "temperature": TEMPERATURE, "passe": PASSE,
               "api": "xAI https://api.x.ai/v1 (response_format json_object)",
               "date_debut_utc": datetime.now(timezone.utc).isoformat(),
               "fichier_source": FICHIER_SOURCE.name, "refus": [], "erreurs_techniques": []}
    tok_in = tok_out = appels = refus = 0
    version_modele = None

    try:
        for t in a_faire:
            if appels >= MAX_APPELS:
                print(f"plafond de {MAX_APPELS} appels atteint")
                break
            appels += 1
            horodatage = datetime.now(timezone.utc).isoformat(timespec="seconds")
            try:
                rep = appeler(client, t["texte"], t["parent"])
            except Exception as e:
                journal["erreurs_techniques"].append({"n": t["n"], "erreur": str(e)})
                print(f"n°{t['n']} : erreur ({e})")
                time.sleep(PAUSE_ENTRE_APPELS)
                continue

            version_modele = getattr(rep, "model", None) or MODELE
            u = getattr(rep, "usage", None)
            ti = getattr(u, "prompt_tokens", 0) or 0
            to = getattr(u, "completion_tokens", 0) or 0
            tok_in += ti; tok_out += to
            texte_rep = rep.choices[0].message.content if rep.choices else ""
            fin = rep.choices[0].finish_reason if rep.choices else "aucun choix"

            d, motif = parser_reponse(texte_rep)
            if d is None:
                refus += 1
                journal["refus"].append({"n": t["n"], "motif": motif, "finish_reason": str(fin),
                                         "brut": (texte_rep or "")[:1000]})
                ecrire_ligne(ws, [t["n"], t["id"], None, None, None, None, f"REFUS ({motif} ; {fin})",
                                  version_modele, horodatage, ti, to], a_reecrire)
                print(f"n°{t['n']} : refus ({motif} ; {fin})")
            else:
                codes = ",".join(str(c) for c in d.get("codes", []) or [])
                ecrire_ligne(ws, [t["n"], t["id"], int(d["classe"]), d.get("incertain", ""), codes,
                                  d.get("justification", ""), "", version_modele, horodatage, ti, to], a_reecrire)
                print(f"n°{t['n']} : classe {d['classe']}"
                      + (", incertain" if d.get("incertain") == "oui" else "")
                      + (f", codes [{codes}]" if codes else ""))

            if appels % 10 == 0:
                wb.save(FICHIER_SORTIE)
            time.sleep(PAUSE_ENTRE_APPELS)
    finally:
        wb.save(FICHIER_SORTIE)
        cout = tok_in / 1e6 * PRIX_ENTREE + tok_out / 1e6 * PRIX_SORTIE
        journal.update({"date_fin_utc": datetime.now(timezone.utc).isoformat(),
                        "version_modele_api": version_modele, "appels": appels,
                        "nb_refus": refus, "tokens_entree": tok_in, "tokens_sortie": tok_out,
                        "cout_estime_usd": round(cout, 4)})
        sessions = []
        if os.path.exists(FICHIER_JOURNAL):
            try:
                with open(FICHIER_JOURNAL, encoding="utf-8") as f:
                    ancien = json.load(f)
                sessions = ancien.get("sessions", [ancien])
            except Exception:
                pass
        sessions.append(journal)
        with open(FICHIER_JOURNAL, "w", encoding="utf-8") as f:
            json.dump({"sessions": sessions}, f, ensure_ascii=False, indent=2)
        print(f"\n{appels} appels, {refus} refus, {cout:.3f} $")
        print(f"{FICHIER_SORTIE.name}\n{FICHIER_JOURNAL.name}")

if __name__ == "__main__":
    main()
