# -*- coding: utf-8 -*-
"""Annotation du corpus lexical par Gemini : 1 appel par tweet, T=0, JSON forcé.
Refus consignés sans relance, reprise en place, journal des tokens et du coût.
"""

import json, os, sys, time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from google import genai
from google.genai import types
from google.genai import errors as genai_errors


MODELE = "gemini-3.5-flash"
TEMPERATURE = 0.0

DEPOT = Path(__file__).resolve().parent.parent
SOURCES = Path(os.environ.get("CORPUS_SOURCES", DEPOT.parent / "donnees_sources"))

FICHIER_SOURCE = SOURCES / "TEXTES échantillon 200 (ne pas utiliser les colonnes annotation).xlsx"
FEUILLE_SOURCE = "Annotation"

PASSE = 1
LIGNES = None
MAX_APPELS = 250
PAUSE_ENTRE_APPELS = 1.0
FICHIER_SORTIE = SOURCES / f"annotations_gemini_lexical_passe{PASSE}.xlsx"
FICHIER_JOURNAL = SOURCES / f"journal_gemini_lexical_passe{PASSE}.json"


PRIX_ENTREE, PRIX_SORTIE = 1.50, 9.00


PROMPT = (DEPOT / "consignes" / "lexical.md").read_text(encoding="utf-8")



def charger_tweets():
    wb = openpyxl.load_workbook(FICHIER_SOURCE, read_only=True, data_only=True)
    ws = wb[FEUILLE_SOURCE]
    tweets = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        n = int(row[0])
        tweets.append({"n": n, "id": str(row[1]), "texte": str(row[2] or ""),
                       "parent": "" if row[3] in (None, "nan") else str(row[3])})
    wb.close()
    return tweets

def preparer_sortie():
    a_reecrire = {}
    if os.path.exists(FICHIER_SORTIE):
        wb = openpyxl.load_workbook(FICHIER_SORTIE)
        ws = wb.active
        faits = set()
        for idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if r[0] is None:
                continue
            n = int(r[0])
            if r[2] is not None:          
                faits.add(n)
            else:                          
                a_reecrire[n] = idx
        print(f"reprise : {len(faits)} faites, {len(a_reecrire)} à refaire")
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Annotation LLM"
        ws.append(["n°", "id_tweet", "classe", "incertain", "codes", "justification",
                   "refus", "modele_version", "horodatage_utc", "tokens_entree", "tokens_sortie"])
        faits = set()
    return wb, ws, faits, a_reecrire

def ecrire_ligne(ws, valeurs, a_reecrire):
    n = valeurs[0]
    if n in a_reecrire:
        for j, v in enumerate(valeurs, start=1):
            ws.cell(row=a_reecrire[n], column=j, value=v)
    else:
        ws.append(valeurs)

def parser_reponse(texte):
    if not texte or not texte.strip():
        return None, "reponse vide"
    t = texte.strip()
    d = None
    try:
        d = json.loads(t)
    except json.JSONDecodeError:
        deb = t.find("{")
        if deb == -1:
            return None, "pas de JSON dans la reponse"
        frag = t[deb:].rstrip()
        for _ in range(10):
            try:
                d = json.loads(frag)
                break
            except json.JSONDecodeError:
                ouv, fer = frag.count("{"), frag.count("}")
                if fer > ouv and frag.endswith("}"):
                    frag = frag[:-1].rstrip()     
                elif ouv > fer:
                    frag = frag + "}" * (ouv - fer)  
                else:
                    try:
                        d = json.loads(frag + '"}') 
                    except json.JSONDecodeError:
                        return None, "JSON irreparable"
                    break
        if d is None:
            return None, "JSON irreparable"
    if not isinstance(d, dict):
        return None, "JSON inattendu (pas un objet)"
    if d.get("refus") is True:
        return None, "refus explicite du modele"
    if d.get("classe") not in (0, 1, 2):
        return None, f"classe invalide : {d.get('classe')!r}"
    return d, None


FILTRES_SECURITE = [
    types.SafetySetting(category=c, threshold="BLOCK_NONE")
    for c in ("HARM_CATEGORY_HARASSMENT", "HARM_CATEGORY_HATE_SPEECH",
              "HARM_CATEGORY_SEXUALLY_EXPLICIT", "HARM_CATEGORY_DANGEROUS_CONTENT")
]

def appeler(client, texte, parent):
    prompt = PROMPT.replace("{{TEXTE}}", texte).replace("{{TEXTE_PARENT}}", parent)
    for tentative in (1, 2, 3): 
        try:
            return client.models.generate_content(
                model=MODELE, contents=prompt,
                config=types.GenerateContentConfig(
                    temperature=TEMPERATURE,
                    response_mime_type="application/json",
                    safety_settings=FILTRES_SECURITE))
        except genai_errors.APIError as e:
            if e.code in (429, 500, 503) and tentative < 3:
                attente = 15 * tentative
                print(f"  {e.code} — nouvelle tentative dans {attente}s")
                time.sleep(attente)
            else:
                raise



def main():
    cle = os.environ.get("GEMINI_API_KEY", "")
    if not cle:
        sys.exit("Définissez GEMINI_API_KEY.")
    client = genai.Client(api_key=cle)

    tweets = charger_tweets()
    if LIGNES is not None:
        tweets = [t for t in tweets if t["n"] in set(LIGNES)]
    wb, ws, faits, a_reecrire = preparer_sortie()
    a_faire = [t for t in tweets if t["n"] not in faits]
    print(f"{len(a_faire)} tweets à annoter — {MODELE}, T={TEMPERATURE}, passe {PASSE}")

    journal = {"modele_demande": MODELE, "temperature": TEMPERATURE, "passe": PASSE,
               "filtres_securite": "BLOCK_NONE (toutes categories) — decision methodologique, cf. ch. 5",
               "date_debut_utc": datetime.now(timezone.utc).isoformat(),
               "fichier_source": FICHIER_SOURCE.name, "refus": [], "erreurs_techniques": []}
    tok_in = tok_out = appels = refus = 0
    version_modele = None

    try:
        for t in a_faire:
            if appels >= MAX_APPELS:
                print(f"plafond de {MAX_APPELS} appels atteint")
                break
            appels += 1
            horodatage = datetime.now(timezone.utc).isoformat(timespec="seconds")
            try:
                rep = appeler(client, t["texte"], t["parent"])
            except Exception as e:
                journal["erreurs_techniques"].append({"n": t["n"], "erreur": str(e)})
                print(f"n°{t['n']} : erreur ({e})")
                time.sleep(PAUSE_ENTRE_APPELS)
                continue

            version_modele = getattr(rep, "model_version", None) or MODELE
            u = rep.usage_metadata
            ti = getattr(u, "prompt_token_count", 0) or 0
            to = getattr(u, "candidates_token_count", 0) or 0
            tok_in += ti; tok_out += to

            d, motif = parser_reponse(rep.text)
            if d is None:
                refus += 1
               
                fin = None
                try:
                    if rep.candidates:
                        fin = str(rep.candidates[0].finish_reason)
                    if getattr(rep, "prompt_feedback", None) and rep.prompt_feedback.block_reason:
                        fin = f"{fin} / bloque_entree={rep.prompt_feedback.block_reason}"
                except Exception:
                    pass
                journal["refus"].append({"n": t["n"], "motif": motif, "finish_reason": fin,
                                         "brut": (rep.text or "")[:1000]})
                ecrire_ligne(ws, [t["n"], t["id"], None, None, None, None, f"REFUS ({motif} ; {fin})",
                                  version_modele, horodatage, ti, to], a_reecrire)
                print(f"n°{t['n']} : refus ({motif} ; {fin})")
            else:
                codes = ",".join(str(c) for c in d.get("codes", []) or [])
                ecrire_ligne(ws, [t["n"], t["id"], int(d["classe"]), d.get("incertain", ""), codes,
                                  d.get("justification", ""), "", version_modele, horodatage, ti, to], a_reecrire)
                print(f"n°{t['n']} : classe {d['classe']}"
                      + (f", incertain" if d.get("incertain") == "oui" else "")
                      + (f", codes [{codes}]" if codes else ""))

            if appels % 10 == 0:
                wb.save(FICHIER_SORTIE)
            time.sleep(PAUSE_ENTRE_APPELS)
    finally:
        wb.save(FICHIER_SORTIE)
        cout = tok_in / 1e6 * PRIX_ENTREE + tok_out / 1e6 * PRIX_SORTIE
        journal.update({"date_fin_utc": datetime.now(timezone.utc).isoformat(),
                        "version_modele_api": version_modele, "appels": appels,
                        "nb_refus": refus, "tokens_entree": tok_in, "tokens_sortie": tok_out,
                        "cout_estime_usd": round(cout, 4)})
        
        sessions = []
        if os.path.exists(FICHIER_JOURNAL):
            try:
                with open(FICHIER_JOURNAL, encoding="utf-8") as f:
                    ancien = json.load(f)
                sessions = ancien.get("sessions", [ancien])
            except Exception:
                pass
        sessions.append(journal)
        with open(FICHIER_JOURNAL, "w", encoding="utf-8") as f:
            json.dump({"sessions": sessions}, f, ensure_ascii=False, indent=2)
        print(f"\n{appels} appels, {refus} refus, {cout:.3f} $")
        print(f"{FICHIER_SORTIE.name}\n{FICHIER_JOURNAL.name}")

if __name__ == "__main__":
    main()
